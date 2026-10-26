import asyncio
import io
import json
import csv
import shutil
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field

from app import config, reports_store, store
from app.config import DATA, MAX_UPLOAD, MAX_PAGES, DETECTOR, RECOGNIZER
from app.services.documents import decode_pages, inference_image
from app.services.ocr_engine import OCREngine
from app.upload_limit import UploadBodyLimit
from app.services.layout import order_document, text_for_page
from app.services.aavin_parser import AavinAbstractParser
from app.services.email_service import AavinReportEmailService
from app.tables import router as tables_router
from app.reports_api import router as reports_router

engine = OCREngine()
executor = None
admission = asyncio.Semaphore(1)


@asynccontextmanager
async def lifespan(app):
    global executor
    store.initialize()
    reports_store.initialize_reports_tables()
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix='ocr')
    yield
    executor.shutdown(wait=True, cancel_futures=True)


app = FastAPI(title='Bilingual OCR Studio', lifespan=lifespan)
app.include_router(tables_router)
app.include_router(reports_router)
app.add_middleware(UploadBodyLimit)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=['127.0.0.1', 'localhost', 'testserver'])


@app.middleware('http')
async def local_origin(request, call_next):
    # Reject cross-site writes even when they use multipart/simple requests.
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        if origin and origin not in ('http://127.0.0.1:5173', 'http://localhost:5173', 'http://127.0.0.1:8000', 'http://localhost:8000'):
            return Response('Untrusted origin', status_code=403)
    response = await call_next(request)
    response.headers['Cache-Control'] = 'no-store'
    return response


def document(identifier):
    doc = store.get(identifier)
    if doc is None:
        raise HTTPException(404, 'Document not found')
    return order_document(doc)


def process(identifier):
    doc = store.get(identifier)
    started = time.perf_counter()
    doc.update(status='processing', error=None)
    store.save(doc)
    try:
        lines = []
        for page in doc['pages']:
            image, sx, sy = inference_image(DATA / identifier / page['image'])
            lines.extend(engine.recognize(image, page['index'], sx, sy))
            doc['completed_pages'] += 1
            store.save(doc)
        doc.update(status='complete', lines=lines, warning=None if lines else 'No text detected. Try a clearer image.')
        # Automated Aavin Production Abstract extraction and persistence from live OCR
        if lines:
            full_text = " ".join((l.get('corrected_text') or l.get('raw_text') or '') for l in lines)
            if AavinAbstractParser.is_aavin_abstract(full_text):
                try:
                    config.reload_env()
                    parsed_report = AavinAbstractParser.parse_aavin_abstract(
                        raw_text=full_text,
                        ocr_created_at=doc.get('created_at'),
                        document_id=identifier
                    )
                    if parsed_report:
                        items = parsed_report.pop('items')
                        saved = reports_store.save_or_reconcile_report(parsed_report, items, user='live_ocr_pipeline')
                        doc['aavin_report_id'] = saved['id']
                        if saved['validation_status'] == 'VALIDATED' and config.REPORT_EMAIL_ENABLED:
                            AavinReportEmailService.send_report_email(saved['id'], user='auto_pipeline', force=True)
                except Exception:
                    pass
    except Exception as exc:
        # Never persist/log arbitrary exception strings from OCR libraries or documents.
        reason = 'Local models missing; run scripts/cache_models.py.' if isinstance(exc, RuntimeError) and 'models missing' in str(exc) else 'OCR initialization or inference failed. Run scripts/verify_ocr.py for a synthetic diagnostic.'
        doc.update(status='failed', error=reason)
    finally:
        doc['processing_seconds'] = round(time.perf_counter() - started, 3)
        store.save(doc)


@app.get('/api/health')
def health():
    return dict(status='ok', engine_loaded=engine._pipeline is not None, detector=DETECTOR, recognizer=RECOGNIZER,
                max_upload_mb=MAX_UPLOAD // 1024 ** 2, max_pages=MAX_PAGES,
                reports_subsystem='ready', reporting=config.safe_report_config())


@app.get('/api/documents')
def list_documents():
    return store.summaries()


@app.post('/api/documents', status_code=202)
async def upload(file: UploadFile = File(...)):
    async with admission:
        if any(d['status'] in ('queued', 'processing') for d in store.summaries()):
            raise HTTPException(409, 'An OCR job is already running. Wait for it to finish.')
        extension = Path(file.filename or '').suffix.lower()
        if extension not in ('.jpg', '.jpeg', '.png', '.pdf'):
            raise HTTPException(415, 'Upload PNG, JPG, JPEG or PDF')
        identifier = str(uuid.uuid4())
        folder = DATA / identifier
        folder.mkdir(parents=True)
        source = folder / ('source' + extension)
        try:
            total = 0
            with source.open('wb') as target:
                while chunk := await file.read(1024 * 1024):
                    total += len(chunk)
                    if total > MAX_UPLOAD:
                        raise HTTPException(413, 'Upload exceeds size limit')
                    target.write(chunk)
            pages = await asyncio.to_thread(decode_pages, source, folder)
            doc = dict(id=identifier, name=Path(file.filename).name, created_at=datetime.now(timezone.utc).isoformat(),
                       status='queued', pages=pages, lines=[], history=[], completed_pages=0, error=None)
            store.save(doc)
        except HTTPException:
            shutil.rmtree(folder)
            raise
        except Exception:
            shutil.rmtree(folder)
            raise HTTPException(422, 'Cannot decode document. Check format, page count, pixel limit or PDF password.') from None
        finally:
            await file.close()
        executor.submit(process, identifier)
        return doc


@app.get('/api/documents/{identifier}')
def get_document(identifier: str):
    return document(identifier)


@app.get('/api/documents/{identifier}/pages/{index}')
def page_image(identifier: str, index: int):
    doc = document(identifier)
    if index < 0 or index >= len(doc['pages']):
        raise HTTPException(404, 'Page not found')
    return FileResponse(DATA / doc['id'] / doc['pages'][index]['image'], media_type='image/png', headers={'Cache-Control': 'no-store'})


class Correction(BaseModel):
    text: str = Field(max_length=10000)
    reviewed: bool


@app.patch('/api/documents/{identifier}/lines/{line_id}')
def correction(identifier: str, line_id: str, change: Correction):
    if document(identifier)['status'] != 'complete':
        raise HTTPException(409, 'Wait for OCR to finish')
    try:
        return order_document(store.correct(identifier, line_id, change.text, change.reviewed))
    except KeyError:
        raise HTTPException(404, 'Line not found') from None


@app.get('/api/documents/{identifier}/export/{format}')
def export(identifier: str, format: str, metadata: bool = True):
    doc = document(identifier)
    if doc['status'] != 'complete':
        raise HTTPException(409, 'Wait for OCR to finish')
    rows = [[line['page_index'] + 1, line['reading_order'] + 1,
             line['corrected_text'] if line['corrected_text'] is not None else line['raw_text'],
             line['raw_text'], line['confidence'], line['reviewed']] for line in doc['lines']]
    headers = ['Page', 'Line', 'Text', 'Raw text', 'Confidence', 'Reviewed']
    if format == 'json':
        content = json.dumps(doc if metadata else [dict(zip(headers, r)) for r in rows], ensure_ascii=False, indent=2).encode('utf-8')
        mime = 'application/json'
    elif format == 'txt':
        content = '\n\f\n'.join(text_for_page([l for l in doc['lines'] if l['page_index'] == p['index']]) for p in doc['pages']).encode('utf-8')
        mime = 'text/plain; charset=utf-8'
    elif format == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        # Prevent spreadsheet formula execution; raw exact strings remain in JSON/TXT.
        writer.writerows([["'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else v for v in row] for row in rows])
        content = output.getvalue().encode('utf-8-sig')
        mime = 'text/csv; charset=utf-8'
    elif format == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Text lines'
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['C'].width = 60
        sheet.column_dimensions['D'].width = 60
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        raise HTTPException(400, 'Unknown export format')
    return Response(content, media_type=mime, headers={'Content-Disposition': f'attachment; filename="ocr-{identifier}.{format}"', 'Cache-Control': 'no-store'})
