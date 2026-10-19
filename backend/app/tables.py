import csv
import io
import json
from datetime import datetime, timezone
from typing import Literal

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field, model_validator

from app import store
from app.config import DATA
from app.services.table_parser import materialize, suggest_layout

router = APIRouter(prefix='/api/documents/{identifier}/tables/{page_index}')


def get_page(identifier, page_index):
    doc = store.get(identifier)
    if not doc:
        raise HTTPException(404, 'Document not found')
    if doc['status'] != 'complete':
        raise HTTPException(409, 'Wait for OCR to finish')
    if not 0 <= page_index < len(doc['pages']):
        raise HTTPException(404, 'Page not found')
    return doc, doc['pages'][page_index]


class Column(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    kind: Literal['text', 'phone', 'aadhaar', 'pashu_aadhaar'] = 'text'


class Grid(BaseModel):
    x_edges: list[float] = Field(min_length=2, max_length=31)
    y_edges: list[float] = Field(min_length=2, max_length=301)
    columns: list[Column] = Field(min_length=1, max_length=30)
    expected_revision: int = Field(ge=0)
    reset_edits: bool = False

    @model_validator(mode='after')
    def valid_grid(self):
        import math
        for edges in (self.x_edges, self.y_edges):
            if any(not math.isfinite(e) or e < 0 for e in edges) or any(b - a < 1 for a, b in zip(edges, edges[1:])):
                raise ValueError('Boundaries must be finite, increasing, nonnegative and at least one pixel apart')
        if len(self.columns) != len(self.x_edges) - 1:
            raise ValueError('Each column needs two boundaries')
        return self


class CellEdit(BaseModel):
    text: str = Field(max_length=10000)
    reviewed: bool
    expected_revision: int = Field(ge=1)


@router.get('')
def get_table(identifier: str, page_index: int, suggest: bool = False):
    doc, page = get_page(identifier, page_index)
    saved = doc.get('tables', {}).get(str(page_index))
    if saved and not suggest:
        return materialize(saved, doc['lines'])
    draft = suggest_layout(DATA / doc['id'] / page['image'], page, [l for l in doc['lines'] if l['page_index'] == page_index])
    draft['revision'] = saved['revision'] if saved else 0
    return materialize(draft, doc['lines'])


@router.put('')
def save_grid(identifier: str, page_index: int, change: Grid):
    _, page = get_page(identifier, page_index)
    if change.x_edges[-1] > page['width'] or change.y_edges[-1] > page['height']:
        raise HTTPException(422, 'Boundaries exceed page dimensions')

    def operation(doc):
        tables = doc.setdefault('tables', {})
        previous = tables.get(str(page_index))
        revision = previous['revision'] if previous else 0
        if revision != change.expected_revision:
            raise HTTPException(409, 'Table changed in another tab. Reload the table before saving.')
        if previous and previous.get('overrides') and not change.reset_edits:
            raise HTTPException(409, 'Rebuilding the grid clears cell edits. Confirm reset or keep the existing grid.')
        history = previous.get('history', []).copy() if previous else []
        if previous:
            history.append(dict(action='rebuild', at=datetime.now(timezone.utc).isoformat(),
                previous={k: v for k, v in previous.items() if k != 'history'}))
        table = dict(page_index=page_index, x_edges=change.x_edges, y_edges=change.y_edges,
                     columns=[c.model_dump() for c in change.columns], revision=revision + 1, overrides={}, history=history)
        tables[str(page_index)] = table
        return materialize(table, doc['lines'])
    return store.mutate(identifier, operation)


@router.patch('/cells/{row}/{column}')
def edit_cell(identifier: str, page_index: int, row: int, column: int, change: CellEdit):
    get_page(identifier, page_index)

    def operation(doc):
        table = doc.get('tables', {}).get(str(page_index))
        if table is None:
            raise HTTPException(409, 'Save the table layout before editing cells')
        if table['revision'] != change.expected_revision:
            raise HTTPException(409, 'Table changed in another tab. Reload the table before saving.')
        if not (0 <= row < len(table['y_edges']) - 1 and 0 <= column < len(table['columns'])):
            raise HTTPException(404, 'Cell not found')
        key = f'{row}:{column}'
        before = materialize(table, doc['lines'])['rows'][row][column]
        table.setdefault('history', []).append(dict(action='cell', row=row, column=column,
            before=before['text'], after=change.text, previous_reviewed=before['reviewed'],
            reviewed=change.reviewed, at=datetime.now(timezone.utc).isoformat()))
        table['overrides'][key] = dict(text=change.text, reviewed=change.reviewed)
        table['revision'] += 1
        return materialize(table, doc['lines'])
    return store.mutate(identifier, operation)


@router.get('/export/{format}')
def export_table(identifier: str, page_index: int, format: str, metadata: bool = True):
    doc, _ = get_page(identifier, page_index)
    saved = doc.get('tables', {}).get(str(page_index))
    if not saved:
        raise HTTPException(409, 'Save the table layout before exporting')
    table = materialize(saved, doc['lines'])
    headers = [column['name'] for column in table['columns']]
    rows = [[cell['text'] for cell in row] for row in table['rows']]
    if format == 'json':
        payload = {**table, 'source_lines': [l for l in doc['lines'] if l['page_index'] == page_index]} if metadata else {'columns': headers, 'rows': rows}
        content = json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')
        mime = 'application/json'
    elif format == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        for row in [headers] + rows:
            writer.writerow(["'" + v if v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else v for v in row])
        content, mime = output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = f'Page {page_index + 1}'
        for row in [headers] + rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'
                cell.number_format = '@'
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='28755B')
        for i in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 26
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        review = book.create_sheet('Review')
        review.append(['Row', 'Column', 'Reviewed', 'Warnings'])
        for row in table['rows']:
            for cell in row:
                review.append([cell['row'] + 1, headers[cell['column']], cell['reviewed'], '; '.join(cell['warnings'])])
                review.cell(review.max_row, 2).data_type = 's'
        output = io.BytesIO()
        book.save(output)
        content, mime = output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif format == 'txt':
        # Tab-separated rows with quoted multiline fields keep one logical record per row.
        output = io.StringIO(newline='')
        csv.writer(output, delimiter='\t').writerows([headers] + rows)
        content, mime = output.getvalue().encode('utf-8'), 'text/plain; charset=utf-8'
    else:
        raise HTTPException(400, 'Unknown export format')
    return Response(content, media_type=mime, headers={'Content-Disposition':
        f'attachment; filename="table-{identifier}-page-{page_index + 1}.{format}"', 'Cache-Control': 'no-store'})
