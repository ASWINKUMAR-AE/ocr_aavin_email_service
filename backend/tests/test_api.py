import io
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app import main, store


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'DATA', tmp_path)
    monkeypatch.setattr(store, 'DATA', tmp_path)
    # Unit tests isolate HTTP/persistence; verify_ocr.py exercises actual inference.
    monkeypatch.setattr(main.engine, 'recognize', lambda *args: [dict(id='0:0', page_index=0, reading_order=0,
        raw_text='தமிழ் 000000000001', corrected_text=None, confidence=.55, reviewed=False, polygon=[[1,1],[20,1],[20,10],[1,10]])])
    with TestClient(main.app) as client:
        yield client


def upload_complete(client):
    data = io.BytesIO()
    Image.new('RGB', (100,100), 'white').save(data, format='PNG')
    response = client.post('/api/documents', files={'file':('synthetic.png',data.getvalue(),'image/png')})
    assert response.status_code == 202
    identifier = response.json()['id']
    for _ in range(100):
        doc = client.get(f'/api/documents/{identifier}').json()
        if doc['status'] == 'complete':
            return doc
        time.sleep(.02)
    pytest.fail('Job did not finish')


def test_upload_correct_history_exports(client):
    doc = upload_complete(client)
    root = f"/api/documents/{doc['id']}"
    response = client.patch(root + '/lines/0:0', json={'text':'தமிழ் 000000000002', 'reviewed':True})
    assert response.status_code == 200
    assert response.json()['lines'][0]['raw_text'] == 'தமிழ் 000000000001'
    assert response.json()['history'][0]['after'] == 'தமிழ் 000000000002'
    assert client.get(root + '/export/txt').content.decode('utf-8') == 'தமிழ் 000000000002'
    assert client.get(root + '/export/csv').content.startswith(b'\xef\xbb\xbf')
    assert client.get(root + '/pages/0').headers['content-type'] == 'image/png'
    assert client.get(root + '/pages/-1').status_code == 404
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(client.get(root + '/export/xlsx').content))
    assert workbook.active['C2'].value == 'தமிழ் 000000000002'
    assert client.get(root + '/export/json').json()['lines'][0]['polygon'] == doc['lines'][0]['polygon']


def test_blank_correction_and_formula_export(client):
    doc = upload_complete(client)
    root = f"/api/documents/{doc['id']}"
    client.patch(root + '/lines/0:0', json={'text':'', 'reviewed':True})
    assert client.get(root + '/export/txt').text == ''
    client.patch(root + '/lines/0:0', json={'text':'=1+1', 'reviewed':False})
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(client.get(root + '/export/xlsx').content))
    assert workbook.active['C2'].data_type == 's'
    assert "'=1+1" in client.get(root + '/export/csv').text


def test_bad_upload_and_cross_origin(client, monkeypatch):
    monkeypatch.setattr(main, 'trusted_origin_set', {'http://localhost:5173'})
    assert client.post('/api/documents', files={'file':('bad.png',b'bad')}).status_code == 422
    assert client.post('/api/documents', files={'file':('bad.exe',b'bad')}).status_code == 415
    assert client.post('/api/documents', headers={'origin':'https://untrusted.example'}, files={'file':('bad.png',b'bad')}).status_code == 403
    monkeypatch.setattr(main, 'MAX_UPLOAD', 2)
    assert client.post('/api/documents', files={'file':('bad.png',b'bad')}).status_code == 413


def test_restart_marks_interrupted_jobs(client):
    store.save(dict(id='test', status='processing'))
    store.initialize()
    assert store.get('test')['status'] == 'failed'


def test_body_rejected_before_multipart_parse(client):
    response = client.post('/api/documents', content=b'x', headers={'content-length':str(100 * 1024 * 1024)})
    assert response.status_code == 413


def test_responses_are_not_browser_cached(client):
    assert client.get('/api/documents').headers['cache-control'] == 'no-store'
