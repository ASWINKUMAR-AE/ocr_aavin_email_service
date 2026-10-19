import io
import json
import csv

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main, store, tables
from app.services.table_parser import materialize, suggest_layout, validate_value
from test_layout import line


def grid():
    return dict(page_index=0, x_edges=[0,100,300], y_edges=[50,150,250],
        columns=[{'name':'Name','kind':'text'},{'name':'Phone','kind':'phone'}],
        revision=1, overrides={}, history=[])


def test_multiline_cells_preserve_raw_empty_merged_and_ditto():
    lines=[line('a',110,60,'first address'), line('b',110,100,'second line'), line('c',110,180,'"')]
    result=materialize(grid(),lines)
    assert result['rows'][0][1]['text'] == 'first address\nsecond line'
    assert result['rows'][0][0]['text'] == result['rows'][1][0]['text'] == ''
    assert result['rows'][1][1]['text'] == '"'
    assert any('ditto' in w for w in result['rows'][1][1]['warnings'])
    assert result['rows'][0][1]['source_ids'] == ['a','b']


def test_number_validation_is_non_mutating():
    assert validate_value('1', 'text') is None
    assert validate_value('01234 56789','phone') is None
    assert validate_value('0000 0000 0001','aadhaar') is None
    assert validate_value('0000 0000 0001','pashu_aadhaar') is None
    assert '10 digits' in validate_value('01234O6789','phone')
    assert validate_value('', 'phone') == 'Missing value'


def test_suggestion_on_synthetic_ruled_table(tmp_path):
    image=np.full((600,800,3),255,np.uint8)
    for x in [20,200,400,780]:
        cv2.line(image,(x,30),(x,580),(150,30,60),3)
    for y in [30,100,220,340,460,580]:
        cv2.line(image,(20,y),(780,y),(150,30,60),3)
    path=tmp_path/'grid.png'
    cv2.imwrite(str(path),image)
    lines=[line(str(i),450,140+i*120,'000000000001') for i in range(4)]
    table=suggest_layout(path,{'width':800,'height':600,'index':0},lines)
    assert len(table['columns']) == 3
    assert len(table['y_edges']) - 1 == 4
    assert 95 <= table['y_edges'][0] <= 105


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA',tmp_path)
    monkeypatch.setattr(tables,'DATA',tmp_path)
    with TestClient(main.app) as client:
        store.save(dict(id='synthetic',name='synthetic.png',status='complete',created_at='synthetic',
            pages=[{'index':0,'width':300,'height':250,'image':'page-0.png'}],
            lines=[line('name',10,80,'தமிழ்'),line('number',110,80,'0123456789')], history=[]))
        yield client


ROOT='/api/documents/synthetic/tables/0'


def save_grid(client, revision=0, reset=False):
    g=grid()
    return client.put(ROOT,json={k:g[k] for k in ('x_edges','y_edges','columns')} |
        {'expected_revision':revision,'reset_edits':reset})


def test_table_save_edit_reload_and_exports(client):
    assert save_grid(client).status_code == 200
    response=client.patch(ROOT+'/cells/0/1',json={'text':'0000000001','reviewed':True,'expected_revision':1})
    assert response.status_code == 200
    saved=client.get(ROOT).json()
    assert saved['rows'][0][1]['text']=='0000000001'
    assert saved['rows'][0][1]['raw_text']=='0123456789'
    assert saved['history'][0]['before']=='0123456789'
    assert saved['rows'][0][1]['reviewed']
    assert store.get('synthetic')['lines'][1]['raw_text']=='0123456789'
    csv_response=client.get(ROOT+'/export/csv')
    assert csv_response.content.startswith(b'\xef\xbb\xbf')
    rows=list(csv.reader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
    assert rows[1]==['தமிழ்','0000000001']
    workbook=load_workbook(io.BytesIO(client.get(ROOT+'/export/xlsx').content))
    assert workbook.active['B2'].value=='0000000001'
    assert workbook.active['B2'].data_type=='s'
    assert 'Review' in workbook.sheetnames
    assert client.get(ROOT+'/export/json').json()['source_lines'][0]['id']=='name'
    assert 'source_lines' not in client.get(ROOT+'/export/json?metadata=false').json()
    assert '0000000001' in client.get(ROOT+'/export/txt').text


def test_conflicts_reset_archive_and_blank_correction(client):
    save_grid(client)
    assert client.patch(ROOT+'/cells/0/1',json={'text':'','reviewed':True,'expected_revision':1}).status_code==200
    assert client.get(ROOT).json()['rows'][0][1]['text']==''
    assert client.patch(ROOT+'/cells/0/1',json={'text':'x','reviewed':False,'expected_revision':1}).status_code==409
    assert save_grid(client,2).status_code==409
    assert save_grid(client,2,True).status_code==200
    archive=client.get(ROOT).json()['history'][-1]
    assert archive['previous']['overrides']['0:1']['text']==''


def test_reject_invalid_grid_and_cell_indices(client):
    data={k:grid()[k] for k in ('x_edges','y_edges','columns')} | {'expected_revision':0}
    assert client.put(ROOT,json={**data,'x_edges':[0,300,100]}).status_code==422
    assert client.put(ROOT,json={**data,'x_edges':[0,100,301]}).status_code==422
    assert client.get('/api/documents/synthetic/tables/-1').status_code==404
    assert client.get(ROOT+'/export/csv').status_code==409
    save_grid(client)
    assert client.patch(ROOT+'/cells/-1/0',json={'text':'a','reviewed':False,'expected_revision':1}).status_code==404


def test_formula_like_cells_export_as_text(client):
    save_grid(client)
    client.patch(ROOT+'/cells/0/0',json={'text':'=1+1','reviewed':False,'expected_revision':1})
    book=load_workbook(io.BytesIO(client.get(ROOT+'/export/xlsx').content))
    assert book.active['A2'].value=='=1+1'
    assert book.active['A2'].data_type=='s'
    assert "'=1+1" in client.get(ROOT+'/export/csv').text
